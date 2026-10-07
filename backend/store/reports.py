from io import BytesIO
from xml.sax.saxutils import escape
from django.http import HttpResponse
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer


def export_report(products, kind, output, request):
    title = 'Productos con Stock Bajo' if kind == 'low' else 'Stock por Tienda'
    generated = timezone.localtime().strftime('%d/%m/%Y %H:%M:%S (Lima)')
    headers = ['Tienda', 'SKU', 'Producto', 'Categoría', 'Stock', 'Stock mínimo']
    rows = [[p.store.name, p.sku, p.name, p.category, p.stock, p.minimum_stock] for p in products]
    filters = ' · '.join(f'{label}: {request.query_params[key]}' for key, label in
        [('store', 'Tienda ID'), ('search', 'Búsqueda'), ('category', 'Categoría'), ('low_stock', 'Solo stock bajo')]
        if request.query_params.get(key)) or 'Todos los productos del alcance autorizado'
    buffer = BytesIO()
    if output == 'xlsx':
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = 'Inventario'
        sheet.append([title])
        sheet.append(['Generado: ' + generated])
        sheet.append([filters])
        sheet.append(headers)
        for row in rows:
            sheet.append(row)
        if not rows:
            sheet.append(['Sin resultados para los filtros seleccionados.'])
        for row in sheet:
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = 's'
                cell.alignment = Alignment(vertical='top', wrap_text=True)
        for cell in sheet[4]:
            cell.fill = PatternFill('solid', fgColor='7C3AED')
            cell.font = Font(color='FFFFFF', bold=True)
        sheet.row_dimensions[1].height = 28
        sheet['A1'].font = Font(size=18, bold=True, color='7C3AED')
        for row in range(1, 4):
            sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
        for i, width in enumerate([26, 22, 42, 24, 14, 18], 1):
            sheet.column_dimensions[get_column_letter(i)].width = width
        sheet.freeze_panes = 'A5'
        sheet.auto_filter.ref = f'A4:F{max(4, sheet.max_row)}'
        workbook.save(buffer)
        content_type = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    else:
        styles = getSampleStyleSheet()
        styles['Normal'].fontSize = 8
        styles['Normal'].leading = 11
        document = SimpleDocTemplate(buffer, pagesize=landscape(A4), rightMargin=30, leftMargin=30, topMargin=30, bottomMargin=35)
        content = [Paragraph('TechStore · ' + title, styles['Title']),
                   Paragraph('Generado: ' + generated, styles['Normal']),
                   Paragraph(escape(filters), styles['Normal']), Spacer(1, 16)]
        if rows:
            data = [[Paragraph(h, styles['Normal']) for h in headers]] + [
                [Paragraph(escape(str(v)), styles['Normal']) for v in row] for row in rows]
            table = Table(data, colWidths=[125, 105, 215, 125, 75, 75], repeatRows=1)
            table.setStyle(TableStyle([('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#DDD3F7')),
                ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F4F1F8')]),
                ('VALIGN', (0, 0), (-1, -1), 'TOP'), ('BOTTOMPADDING', (0, 0), (-1, -1), 9),
                ('TOPPADDING', (0, 0), (-1, -1), 9), ('LINEBELOW', (0, 0), (-1, 0), 1, colors.HexColor('#7C3AED'))]))
            content.append(table)
        else:
            content.append(Paragraph('Sin resultados para los filtros seleccionados.', styles['Normal']))
        def footer(canvas, doc):
            canvas.setFont('Helvetica', 8)
            canvas.drawRightString(810, 18, f'TechStore · Página {doc.page}')
        document.build(content, onFirstPage=footer, onLaterPages=footer)
        content_type = 'application/pdf'
    response = HttpResponse(buffer.getvalue(), content_type=content_type)
    response['Content-Disposition'] = f'attachment; filename="techstore-{kind}.{output}"'
    response['Cache-Control'] = 'no-store'
    return response
